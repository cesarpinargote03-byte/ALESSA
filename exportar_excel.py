import json
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


CARPETA = Path(__file__).resolve().parent
ARCHIVO_ENTRADA = CARPETA / "conversaciones.json"
ARCHIVO_SALIDA = CARPETA / "Correos.xlsx"


def leer_conversaciones(ruta):
    if not ruta.is_file():
        raise FileNotFoundError(f"No se encontró el archivo de entrada: {ruta}")

    with ruta.open("r", encoding="utf-8") as archivo:
        datos = json.load(archivo)

    if isinstance(datos, dict):
        body = datos.get("Body")
        datos = body.get("Conversations") if isinstance(body, dict) else None

    if not isinstance(datos, list):
        raise ValueError(
            "El JSON debe ser una lista de conversaciones o tener la forma "
            "{'Body': {'Conversations': [...]}}."
        )

    for indice, conversacion in enumerate(datos, start=1):
        if not isinstance(conversacion, dict):
            raise ValueError(
                f"La conversación número {indice} no es un objeto JSON."
            )

    return datos


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, (str, int, float, bool)):
        return str(valor)
    return ""


def texto_categorias(valor):
    if isinstance(valor, str):
        return valor
    if isinstance(valor, list):
        return ", ".join(
            categoria for categoria in valor if isinstance(categoria, str)
        )
    return ""


def fecha_excel(valor):
    if not isinstance(valor, str) or not valor:
        return None

    try:
        # Excel no admite datetimes con zona horaria; conserva la hora local recibida.
        return datetime.fromisoformat(valor).replace(tzinfo=None)
    except ValueError:
        return valor


def crear_filas(conversaciones):
    filas = []

    for conversacion in conversaciones:
        remitente = conversacion.get("LastSender", {})
        mailbox = remitente.get("Mailbox", {}) if isinstance(remitente, dict) else {}
        if not isinstance(mailbox, dict):
            mailbox = {}

        participantes = conversacion.get("UniqueSenders", [])
        if not isinstance(participantes, list):
            participantes = []

        try:
            tamano_mb = round(float(conversacion.get("Size", 0)) / (1024 * 1024), 2)
        except (TypeError, ValueError):
            tamano_mb = None

        conv_id = conversacion.get("ConversationId")
        if isinstance(conv_id, dict):
            conv_id_str = conv_id.get("Id", "")
        else:
            conv_id_str = str(conv_id or "")

        filas.append(
            {
                "Asunto": texto(conversacion.get("ConversationTopic")),
                "Fecha última actividad": fecha_excel(
                    conversacion.get("LastDeliveryTime")
                ),
                "Categoría": texto_categorias(conversacion.get("Categories")),
                "Remitente": texto(mailbox.get("Name")),
                "Correo remitente": texto(mailbox.get("EmailAddress")),
                "Contenido / Resumen": texto(conversacion.get("Preview")),
                "Participantes": ", ".join(
                    participante
                    for participante in participantes
                    if isinstance(participante, str)
                ),
                "Tiene adjuntos": bool(conversacion.get("HasAttachments", False)),
                "No leídos": conversacion.get("UnreadCount", 0),
                "Tamaño (MB)": tamano_mb,
                "ID Conversación": conv_id_str,
            }
        )

    return filas


def aplicar_formato(ruta, cantidad_filas):
    libro = load_workbook(ruta)
    hoja = libro.active
    hoja.freeze_panes = "A2"

    if cantidad_filas:
        tabla = Table(displayName="TablaConversaciones", ref=hoja.dimensions)
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        hoja.add_table(tabla)

    anchos = {
        "A": 50,  # Asunto
        "B": 22,  # Fecha última actividad
        "C": 26,  # Categoría
        "D": 30,  # Remitente
        "E": 34,  # Correo remitente
        "F": 65,  # Contenido / Resumen
        "G": 30,  # Participantes
        "H": 15,  # Tiene adjuntos
        "I": 12,  # No leídos
        "J": 14,  # Tamaño (MB)
        "K": 40,  # ID Conversación
    }
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    # Formatear columna B (Fecha última actividad) como fecha y hora en Excel
    for celda in hoja["B"][1:]:
        celda.number_format = "yyyy-mm-dd hh:mm"

    libro.save(ruta)


COLUMNAS = [
    "Asunto",
    "Fecha última actividad",
    "Categoría",
    "Remitente",
    "Correo remitente",
    "Contenido / Resumen",
    "Participantes",
    "Tiene adjuntos",
    "No leídos",
    "Tamaño (MB)",
    "ID Conversación",
]


def guardar_excel(conversaciones, ruta_salida=ARCHIVO_SALIDA, nombre_hoja="Conversaciones"):
    filas = crear_filas(conversaciones)
    pd.DataFrame(filas, columns=COLUMNAS).to_excel(
        ruta_salida,
        index=False,
        sheet_name=nombre_hoja,
    )
    aplicar_formato(ruta_salida, len(filas))
    return ruta_salida


def main():
    conversaciones = leer_conversaciones(ARCHIVO_ENTRADA)
    guardar_excel(conversaciones, ARCHIVO_SALIDA)
    print(f"Exportadas {len(conversaciones)} conversaciones a: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()
